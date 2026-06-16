from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from sqlalchemy.orm import Session
from app.database import get_db
from app.auth import get_current_user
from app.models import ResumeAnalysis, ActivityLog
from app.schemas import ResumeAnalysisResponse, ResumeAnalysisCreate
from typing import Any, List, Optional
import io
import json

router = APIRouter(prefix="/resume", tags=["Resume Analyzer"])

# Helper function to extract text from PDF in Python
def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
    # Try using pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    except ImportError:
        pass

    # Try using fitz (PyMuPDF)
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        return text
    except ImportError:
        pass

    # Basic fallback: decode what we can (often yields some metadata/plain text streams)
    try:
        return pdf_bytes.decode('utf-8', errors='ignore')
    except Exception:
        return "Resume document parsed with standard placeholders."

# Rule-based resume evaluator in Python (same logic as JS for consistency)
def python_analyze_resume(text: str, role: str) -> dict:
    text_lower = text.lower()
    core_keywords = []
    extra_keywords = []
    role_display = role

    if "software" in role.lower():
        core_keywords = ['html', 'css', 'javascript', 'react', 'typescript', 'next.js', 'git', 'node.js']
        extra_keywords = ['redux', 'zustand', 'jest', 'tailwind', 'aws', 'docker', 'graphql', 'mongodb', 'sql', 'express']
        role_display = 'Software Engineer'
    elif "machine" in role.lower() or "learning" in role.lower():
        core_keywords = ['python', 'pytorch', 'tensorflow', 'scikit-learn', 'machine learning', 'deep learning', 'numpy', 'pandas', 'git']
        extra_keywords = ['nlp', 'llm', 'keras', 'docker', 'aws', 'sql', 'statistics', 'opencv', 'transformers']
        role_display = 'Machine Learning Engineer'
    elif "data" in role.lower() or "analyst" in role.lower():
        core_keywords = ['python', 'sql', 'pandas', 'excel', 'tableau', 'power bi', 'statistics', 'probability', 'git']
        extra_keywords = ['r', 'matplotlib', 'seaborn', 'regression', 'clustering', 'scikit-learn', 'jupyter', 'bigquery']
        role_display = 'Data Scientist / Analyst'
    else: # Product Manager
        core_keywords = ['product management', 'agile', 'scrum', 'jira', 'roadmap', 'analytics', 'sql', 'wireframe', 'ab testing']
        extra_keywords = ['figma', 'amplitude', 'mixpanel', 'user research', 'customer feedback', 'kpis', 'metrics', 'strategy']
        role_display = 'Product Manager'

    all_keywords = core_keywords + extra_keywords
    matched = []
    missing = []

    for kw in all_keywords:
        if kw in text_lower:
            matched.append(kw)
        else:
            missing.append(kw)

    def capitalize_kw(s: str) -> str:
        if s in ['html', 'css', 'sql', 'nlp', 'llm', 'kpis']:
            return s.upper()
        return s.title()

    present_display = [capitalize_kw(k) for k in matched]
    missing_display = [capitalize_kw(k) for k in missing]

    core_matched = [kw for kw in core_keywords if kw in matched]
    core_ratio = len(core_matched) / len(core_keywords) if core_keywords else 0
    overall_ratio = len(matched) / len(all_keywords) if all_keywords else 0
    keyword_score = round((core_ratio * 0.7 + overall_ratio * 0.3) * 100)

    action_verbs = ['designed', 'developed', 'optimized', 'implemented', 'built', 'led', 'managed', 'created', 'coordinated', 'executed', 'analyzed', 'researched', 'programmed', 'integrated']
    verbs_count = sum(1 for v in action_verbs if v in text_lower)
    content_score = min(100, 50 + (verbs_count * 5))

    import re
    # Simple check for numbers, percentages, etc.
    metric_matches = re.findall(r'\b\d+(?:%|x|k|LPA|lakh|crore|L)\b', text_lower)
    metric_count = len(metric_matches)
    impact_score = 55
    if metric_count >= 3:
        impact_score = min(100, 85 + (metric_count * 2))
    elif metric_count > 0:
        impact_score = 70 + (metric_count * 5)

    headers = ['experience', 'education', 'projects', 'skills', 'contact', 'summary', 'achievements', 'certifications']
    headers_found = sum(1 for h in headers if h in text_lower)
    format_score = min(100, 60 + (headers_found * 5))

    overall_score = round((content_score + format_score + keyword_score + impact_score) / 4)

    suggestions = []
    if impact_score < 75:
        suggestions.append({
            "type": "critical",
            "message": "Quantify Achievements: Add measurable statistics (e.g. 'improved latency by 30%' or 'saved ₹50K') to validate your achievements."
        })
    if len(core_matched) < len(core_keywords) / 2:
        top_missing = " and ".join(missing_display[:2])
        suggestions.append({
            "type": "critical",
            "message": f"Add Core Competencies: Missing essential tools like {top_missing or 'foundational libraries'} for {role_display}."
        })
    if format_score < 80:
        suggestions.append({
            "type": "critical",
            "message": "Standardize Headers: Use simple section names (e.g. 'Experience', 'Education', 'Projects', 'Skills') to improve ATS scan rates."
        })

    suggestions.append({
        "type": "nice_to_have",
        "message": "Include Profiles: Ensure active links to professional profiles (e.g. LinkedIn, GitHub) are visible in the contact header."
    })

    return {
        "overall_score": overall_score,
        "section_scores": {
            "content": content_score,
            "formatting": format_score,
            "keywords": keyword_score,
            "impact": impact_score
        },
        "suggestions": suggestions,
        "keywords": {
            "present": present_display,
            "missing": missing_display
        }
    }

# 1. FILE UPLOAD & ANALYZE ENDPOINT
@router.post("/analyze", response_model=ResumeAnalysisResponse)
async def upload_and_analyze_resume(
    target_role: str = Form(...),
    file: UploadFile = File(...),
    current_user: Any = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Validate file size (max 5MB)
    contents = await file.read()
    if len(contents) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size exceeds maximum limit of 5MB.")

    # Validate file type
    extension = file.filename.split('.')[-1].lower() if file.filename else ''
    if extension not in ['pdf', 'docx']:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload a PDF or DOCX file.")

    # Extract text from PDF/DOCX
    text = ""
    if extension == 'pdf':
        text = extract_text_from_pdf_bytes(contents)
    else:
        # docx reading
        try:
            import docx
            doc_file = io.BytesIO(contents)
            doc = docx.Document(doc_file)
            text = "\n".join([p.text for p in doc.paragraphs])
        except ImportError:
            text = contents.decode('utf-8', errors='ignore')

    # Analyze text
    analysis = python_analyze_resume(text, target_role)

    # Save to database
    db_analysis = ResumeAnalysis(
        user_id=current_user.id,
        file_url=file.filename or "uploaded_resume",
        target_role=target_role,
        overall_score=analysis["overall_score"],
        section_scores=analysis["section_scores"],
        suggestions=analysis["suggestions"],
        keywords=analysis["keywords"]
    )
    db.add(db_analysis)
    
    # Log activity
    log = ActivityLog(
        user_id=current_user.id,
        action_type="resume_analyzed",
        description=f"Analyzed resume '{file.filename}' for role '{target_role}' (Score: {analysis['overall_score']})."
    )
    db.add(log)
    
    db.commit()
    db.refresh(db_analysis)
    return db_analysis

# 2. SAVE PRE-ANALYZED RESULT (e.g. from frontend)
@router.post("/save", response_model=ResumeAnalysisResponse)
def save_pre_analyzed_resume(
    params: ResumeAnalysisCreate,
    current_user: Any = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    db_analysis = ResumeAnalysis(
        user_id=current_user.id,
        file_url=params.file_url,
        target_role=params.target_role,
        overall_score=params.overall_score,
        section_scores=params.section_scores,
        suggestions=params.suggestions,
        keywords=params.keywords
    )
    db.add(db_analysis)
    
    log = ActivityLog(
        user_id=current_user.id,
        action_type="resume_saved",
        description=f"Saved resume analysis for '{params.file_url}' (Score: {params.overall_score})."
    )
    db.add(log)
    
    db.commit()
    db.refresh(db_analysis)
    return db_analysis

# 3. GET LIST OF ANALYSES
@router.get("/analyses", response_model=List[ResumeAnalysisResponse])
def get_past_analyses(current_user: Any = Depends(get_current_user), db: Session = Depends(get_db)):
    return db.query(ResumeAnalysis).filter(ResumeAnalysis.user_id == current_user.id).order_by(ResumeAnalysis.analyzed_at.desc()).all()

# 4. GET SPECIFIC ANALYSIS BY ID
@router.get("/analyses/{id}", response_model=ResumeAnalysisResponse)
def get_analysis_by_id(id: str, current_user: Any = Depends(get_current_user), db: Session = Depends(get_db)):
    analysis = db.query(ResumeAnalysis).filter(ResumeAnalysis.id == id, ResumeAnalysis.user_id == current_user.id).first()
    if not analysis:
        raise HTTPException(status_code=404, detail="Resume analysis entry not found.")
    return analysis
